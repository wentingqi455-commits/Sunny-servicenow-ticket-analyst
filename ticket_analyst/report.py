"""Build the Word report (and a CSV of the masked tickets)."""

from __future__ import annotations

import csv
import io
from datetime import datetime
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Inches, Pt, RGBColor  # noqa: E402

from .stats import Stats  # noqa: E402

BLUE = "#2F5597"
GREY = "#A6A6A6"
MAX_ROWS = 50


# ---- charts -----------------------------------------------------------------
def _png(fig) -> io.BytesIO:
    buf = io.BytesIO()
    fig.tight_layout()
    fig.savefig(buf, format="png", dpi=150)
    plt.close(fig)
    buf.seek(0)
    return buf


def chart_categories(s: Stats) -> io.BytesIO:
    rows = s.by_category[:10][::-1]
    labels = [r[0] for r in rows]
    y = range(len(rows))
    fig, ax = plt.subplots(figsize=(7, 0.4 * len(rows) + 1.2))
    ax.barh([i + 0.2 for i in y], [r[1] for r in rows], height=0.4, color=BLUE, label="This week")
    ax.barh([i - 0.2 for i in y], [r[2] for r in rows], height=0.4, color=GREY, label="Last week")
    ax.set_yticks(list(y), labels)
    ax.set_xlabel("Tickets")
    ax.legend(loc="lower right", frameon=False)
    ax.spines[["top", "right"]].set_visible(False)
    return _png(fig)


def chart_daily(s: Stats) -> io.BytesIO:
    days = [d[5:] for d, _ in s.by_day]
    fig, ax = plt.subplots(figsize=(7, 2.6))
    ax.bar(days, [n for _, n in s.by_day], color=BLUE)
    ax.set_ylabel("Tickets created")
    ax.spines[["top", "right"]].set_visible(False)
    return _png(fig)


# ---- docx helpers -----------------------------------------------------------
def _table(doc, headers: list[str], rows: list[list[Any]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = "" if v is None else str(v)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
            if widths:
                cell.width = Inches(widths[i])
    doc.add_paragraph()
    return table


def _bullets(doc, items: list[str]):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def _note(doc, text: str):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = True
    run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)


def _change(now: int, before: int) -> str:
    if before == 0:
        return "new" if now else "-"
    pct = (now - before) / before * 100
    return f"{pct:+.0f}%"


def _short(text: str, n: int = 90) -> str:
    text = " ".join((text or "").split())
    return text if len(text) <= n else text[: n - 1] + "…"


# ---- report -----------------------------------------------------------------
def build_report(s: Stats, ai: dict[str, Any] | None, path: Path) -> Path:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    title = doc.add_heading("Weekly IT Ticket Analysis", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph(
        f"Period: {s.period_start:%d %b %Y} to {s.period_end:%d %b %Y} (UTC)    "
        f"Generated: {datetime.now():%d %b %Y %H:%M}"
    )

    # 1. Summary
    doc.add_heading("1. Summary", 1)
    if ai and ai.get("executive_summary"):
        doc.add_paragraph(ai["executive_summary"])
    else:
        _note(doc, "AI analysis was not run, so this report contains statistics only.")
    avg = f"{s.avg_resolution_hours:.1f} h" if s.avg_resolution_hours is not None else "-"
    _table(
        doc,
        ["Measure", "Value"],
        [
            ["Tickets created", s.total],
            ["Previous week", f"{s.previous_total} ({_change(s.total, s.previous_total)} change)"],
            ["Still open", s.open_count],
            ["Resolved / closed", s.resolved_count],
            ["Average time to resolve", avg],
            ["SLA breaches", len(s.sla_breaches)],
            ["Hardware tickets", len(s.hardware)],
        ],
        [2.5, 3.5],
    )

    # 2. Categories and trends
    doc.add_heading("2. Top categories and trends", 1)
    if ai:
        _bullets(doc, ai.get("trend_insights", []))
    if s.by_category:
        doc.add_picture(chart_categories(s), width=Inches(6.3))
    _table(
        doc,
        ["Category", "This week", "Last week", "Change"],
        [[c, a, b, _change(a, b)] for c, a, b in s.by_category[:15]],
        [2.8, 1.1, 1.1, 1.0],
    )
    doc.add_heading("Tickets per day", 2)
    doc.add_picture(chart_daily(s), width=Inches(6.3))
    doc.add_heading("By ticket type and priority", 2)
    _table(doc, ["Type", "Tickets"], s.by_type.most_common(), [3, 1.5])
    _table(doc, ["Priority", "Tickets"], sorted(s.by_priority.items()), [3, 1.5])
    doc.add_heading("Busiest assignment groups", 2)
    _table(doc, ["Assignment group", "Tickets"], s.by_group.most_common(10), [3.5, 1.5])

    # 3. Recurring issues
    doc.add_heading("3. Recurring issues", 1)
    issues = (ai or {}).get("recurring_issues", [])
    if issues:
        _table(
            doc,
            ["Issue", "Tickets", "Likely root cause", "Recommendation"],
            [
                [
                    i.get("title", ""),
                    f"{len(i['ticket_numbers'])}: " + ", ".join(i["ticket_numbers"][:8])
                    + (" …" if len(i["ticket_numbers"]) > 8 else ""),
                    i.get("root_cause", ""),
                    i.get("recommendation", ""),
                ]
                for i in issues
            ],
            [1.5, 1.4, 1.7, 1.9],
        )
    else:
        _note(doc, "No recurring issues identified (or AI analysis was not run).")

    # 4. SLA breaches
    doc.add_heading("4. SLA breaches", 1)
    if ai:
        _bullets(doc, ai.get("sla_insights", []))
    if s.sla_breaches:
        _table(doc, ["SLA", "Breaches"], s.sla_by_name.most_common(10), [4, 1.5])
        _table(doc, ["Assignment group", "Breaches"], s.sla_by_group.most_common(10), [4, 1.5])
        doc.add_heading("Breached tickets", 2)
        _table(
            doc,
            ["Ticket", "Priority", "Group", "SLA", "Short description"],
            [
                [b["number"], b["priority"], b["assignment_group"], b["sla"], _short(b["short_description"], 70)]
                for b in s.sla_breaches[:MAX_ROWS]
            ],
            [1.0, 0.9, 1.3, 1.3, 2.0],
        )
        if len(s.sla_breaches) > MAX_ROWS:
            _note(doc, f"Showing {MAX_ROWS} of {len(s.sla_breaches)}. Full list in the CSV export.")
    else:
        doc.add_paragraph("No SLA breaches this week.")

    # 5. Hardware
    doc.add_heading("5. Hardware issues", 1)
    if ai:
        _bullets(doc, ai.get("hardware_insights", []))
    if s.repeat_barcodes:
        doc.add_heading("Devices with more than one ticket", 2)
        _table(
            doc,
            ["Barcode", "Tickets"],
            [[code, ", ".join(nums)] for code, nums in sorted(s.repeat_barcodes.items(), key=lambda x: -len(x[1]))],
            [2, 4],
        )
    if s.hardware:
        doc.add_heading("All hardware tickets", 2)
        _table(
            doc,
            ["Ticket", "Barcode", "Subcategory", "State", "Short description"],
            [
                [t["number"], ", ".join(t.get("barcodes", [])) or "(not found)", t.get("subcategory", ""),
                 t.get("state", ""), _short(t.get("short_description", ""), 70)]
                for t in s.hardware[:MAX_ROWS]
            ],
            [1.0, 1.3, 1.1, 0.9, 2.2],
        )
        if len(s.hardware) > MAX_ROWS:
            _note(doc, f"Showing {MAX_ROWS} of {len(s.hardware)}. Full list in the CSV export.")
    else:
        doc.add_paragraph("No hardware tickets this week.")

    # 6. Recommendations
    if ai and ai.get("recommendations"):
        doc.add_heading("6. Recommended actions", 1)
        _bullets(doc, ai["recommendations"])

    _note(
        doc,
        "Personal information (names, emails, phone numbers, IP addresses) was removed before "
        "analysis. Statistics are calculated directly from ServiceNow data; narrative text is "
        "AI generated and should be checked.",
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


CSV_FIELDS = [
    "number", "table", "category", "subcategory", "priority", "state", "assignment_group",
    "cmdb_ci", "sys_created_on", "resolved_at", "closed_at", "made_sla", "barcodes",
    "short_description", "description",
]


def write_csv(tickets: list[dict[str, Any]], path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as f:  # utf-8-sig opens cleanly in Excel
        w = csv.DictWriter(f, fieldnames=CSV_FIELDS, extrasaction="ignore")
        w.writeheader()
        for t in tickets:
            w.writerow({**t, "barcodes": " ".join(t.get("barcodes", []))})
    return path
